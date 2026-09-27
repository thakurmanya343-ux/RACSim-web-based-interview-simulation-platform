import sys
import json
import os
import re

def extract_text_from_pdf(file_path):
    import pdfplumber
    text_parts = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            extract = page.extract_text()
            if extract:
                text_parts.append(extract)
    return "\n".join(text_parts)

def extract_text_from_docx(file_path):
    import docx
    doc = docx.Document(file_path)
    text_parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            text_parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                text_parts.append(" | ".join(row_text))
    return "\n".join(text_parts)

def parse_file(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    raw_text = ""
    try:
        if ext == ".pdf":
            raw_text = extract_text_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            raw_text = extract_text_from_docx(file_path)
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                raw_text = f.read()
        else:
            raw_text = f"Unsupported file extension {ext}"
    except Exception as e:
        return {"success": False, "error": str(e), "text": ""}
    
    # Extract candidate email if present
    email_match = re.search(r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+', raw_text)
    email = email_match.group(0) if email_match else ""

    # Extract name candidate heuristic
    lines = [line.strip() for line in raw_text.splitlines() if line.strip()]
    candidate_name = ""
    for line in lines[:5]:
        if len(line) < 55 and not any(char in line for char in "@#{}[]:;=<>") and not re.search(r'\b(resume|cv|curriculum|vitae|page)\b', line, re.I):
            clean_name = line.split('-')[0].split('|')[0].strip()
            if clean_name and len(clean_name) > 2:
                candidate_name = clean_name
                break

    return {
        "success": True,
        "raw_text": raw_text,
        "email": email,
        "candidate_name": candidate_name
    }

if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(json.dumps({"success": False, "error": "No file path provided"}))
        sys.exit(1)

    target_path = sys.argv[1]
    result = parse_file(target_path)
    print(json.dumps(result))
